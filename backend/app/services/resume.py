import io
import re

from app.config import MAX_RESUME_CHARS


class ResumeError(ValueError):
    """Raised when an uploaded resume can't be read."""


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _from_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


# Resume PDFs (LaTeX-style fonts) often come out with private glyph codes instead of ligatures/dashes
_GLYPHS = {"\x15": "-", "\x1c": "fi", "\x1d": "fl", "\x88": "•"}

# Words split by font kerning ("Pro ject", "Educa tion"); a fixed list keeps this safe from over-joining
_SPLIT_WORDS = {
    "Educa tion": "Education", "T echnologies": "Technologies", "T echnical": "Technical", "To ols": "Tools",
    "Pro ject": "Project", "Mo del": "Model", "Ob ject": "Object", "Sub ject": "Subject", "Co de": "Code",
    "Op erations": "Operations", "Bo ot": "Boot", "Develop ed": "Developed", "Do cker": "Docker", "Ful l": "Full",
    "Sub queries": "Subqueries", "Data base": "Database",
}


def clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    for bad, good in _GLYPHS.items():
        text = text.replace(bad, good)
    for bad, good in _SPLIT_WORDS.items():
        text = text.replace(bad, good).replace(bad.lower(), good.lower())
    text = re.sub(r"[\x01-\x08\x0b-\x1f\x7f-\x9f]", " ", text)  # any other stray control codes
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_RESUME_CHARS]


def extract_resume_text(filename: str, data: bytes) -> str:
    """Extract plain text from a PDF / DOCX / TXT resume."""
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            text = _from_pdf(data)
        elif name.endswith(".docx"):
            text = _from_docx(data)
        elif name.endswith((".txt", ".md")):
            text = data.decode("utf-8", errors="ignore")
        else:
            raise ResumeError("Unsupported file type. Upload a PDF, DOCX or TXT resume.")
    except ResumeError:
        raise
    except Exception as e:
        raise ResumeError(f"Could not read the file: {e}") from e

    text = clean_text(text)
    if not text:
        raise ResumeError("No text found in the file (scanned/image PDFs are not supported).")
    return text
